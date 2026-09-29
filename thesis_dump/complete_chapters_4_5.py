from docx import Document
from docx.text.paragraph import Paragraph
from pathlib import Path
from docx.shared import Inches

src=Path(r'D:\Baretta\Skripsi\program\sora_academic\thesis_dump\Skripsi_123220204_BarettaP_format_template.docx')
out=Path(r'D:\Baretta\Skripsi\program\sora_academic\thesis_dump\Skripsi_123220204_BarettaP_final_bab4_bab5.docx')
doc=Document(src)
paras={p.text.strip():p for p in doc.paragraphs if p.text.strip()}
for p in doc.paragraphs:
    if p.style.name=='Heading 1' and p.text.strip()=='HASIL DAN PEMBAHASAN':
        p.text='HASIL PENGUJIAN DAN PEMBAHASAN'

def get(starts):
    matches=[p for p in doc.paragraphs if p.text.strip().startswith(starts)]
    if len(matches)!=1: raise ValueError((starts,len(matches)))
    return matches[0]

def setp(starts,text):
    p=get(starts); style=p.style
    p.clear(); p.style=style; p.add_run(text.replace('*','').replace('`',''))
    return p

def after(p,text,style='Normal'):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    e=OxmlElement('w:p'); p._p.addnext(e)
    q=Paragraph(e,p._parent); q.style=style; q.add_run(text.replace('*','').replace('`','')); return q

def before(p,text,style='Normal'):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph
    e=OxmlElement('w:p'); p._p.addprevious(e)
    q=Paragraph(e,p._parent); q.style=style; q.add_run(text.replace('*','').replace('`','')); return q

setp('Bab ini menyajikan hasil eksperimen SORA',
     'Bab ini menyajikan hasil pengujian sistem SORA (*Sleep Optimization and Rest Arrangement*) pada data final, kemudian menjelaskan makna dan batas interpretasi setiap metrik. Dataset pengujian memuat 63 baris dengan 63 responden unik. Setiap responden diproses untuk tujuh hari secara berurutan, sehingga terdapat 441 evaluasi harian. Jadwal kerja dan kuliah diperlakukan sebagai aktivitas wajib, sedangkan jadwal tidur dari kuesioner menjadi baseline observasi. Semua angka pada bab ini merupakan hasil evaluasi komputasional dengan seed 42; skor fitness adalah skor fungsi tujuan yang tidak memiliki satuan dan bukan persentase akurasi atau ukuran klinis.')

setp('Kalibrasi grid search dilakukan terpisah',
     'Kalibrasi parameter dilakukan terpisah dari eksperimen utama menggunakan *grid search*. Tiga parameter yang diuji ialah ukuran populasi 50 dan 100, laju mutasi 0,01; 0,05; dan 0,10, serta laju *crossover* 0,70; 0,80; dan 0,90. Kombinasi tersebut membentuk 18 konfigurasi. Setiap konfigurasi dinilai pada lima subjek dataset kalibrasi dan tiga hari representatif (Senin, Rabu, dan Sabtu). Nilai pada Tabel 4.2 merupakan rata-rata fitness kalibrasi untuk masing-masing kombinasi; skor yang lebih tinggi dipilih sebagai hasil terbaik menurut fungsi tujuan yang digunakan.')

setp('Kombinasi dengan skor kalibrasi tertinggi',
     'Skor tertinggi pada kalibrasi adalah 2,917, yang diperoleh oleh ukuran populasi 100, mutasi 0,10, dan *crossover* 0,80. Nilai kalibrasi yang negatif pada beberapa kombinasi tetap mungkin terjadi karena fitness menjumlahkan manfaat tidur dengan penalti durasi, fragmentasi tidur, *nap*, dan regularitas. Nilai negatif bukan berarti data hilang; nilainya menunjukkan bahwa penalti pada fungsi tujuan lebih besar daripada komponen manfaat pada kasus yang diuji. Namun, hasil kalibrasi ini tidak digunakan sebagai konfigurasi eksperimen utama: metadata final mencatat `parameter_source=default_fallback`, ukuran populasi 50, mutasi 0,10, *crossover* 0,80, dan seed 42. Karena proses utama dijalankan dengan opsi *skip-calibration*, angka 2,917 harus dibaca sebagai hasil eksperimen kalibrasi terpisah, bukan hasil dari konfigurasi yang dipakai untuk menghasilkan 441 jadwal.')

p=get('Skor tertinggi pada kalibrasi')
after(p,'Grid search membentuk 18 kombinasi dari 2 ukuran populasi, 3 nilai mutasi, dan 3 nilai crossover. Pada setiap konfigurasi, lima subjek kalibrasi dinilai pada tiga hari representatif (Senin, Rabu, Sabtu); implementasi merata-ratakan skor hari dalam subjek terlebih dahulu, lalu merata-ratakan lima subjek. Nilai rata-rata 18 konfigurasi berada pada rentang −263,881 sampai 2,917. Enam konfigurasi memiliki rata-rata negatif sekitar −263,8, sedangkan dua belas konfigurasi lainnya memiliki rata-rata positif sekitar 2,84–2,92. Berkas hasil kalibrasi hanya menyimpan rata-rata per konfigurasi, bukan rincian fitness dan penalti per subjek-hari; karena itu penyebab spesifik skor negatif tiap konfigurasi tidak dapat dipastikan dari artefak final dan tidak diatribusikan ke satu aturan tertentu. Selain itu, kode menetapkan seed berbeda untuk indeks konfigurasi yang berbeda, sehingga selisih antarkonfigurasi juga dapat memuat variasi inisialisasi acak, bukan hanya pengaruh parameter.')

setp('Pengujian mingguan dijalankan pada seluruh 63 responden',
     'Pengujian mingguan memakai satu baris kuesioner untuk setiap responden dan memproses hari Senin sampai Minggu secara berurutan. Seluruh 63 responden menghasilkan tujuh jadwal; 441 dari 441 hari berstatus optimized. Status tersebut berarti proses optimasi menghasilkan keluaran harian, bukan bukti bahwa setiap jadwal sudah ideal secara medis. Pada setiap hari, jadwal kerja dan kuliah serta slot transisinya dipertahankan sebagai kendala; sistem mengoptimasi aktivitas tidur pada slot lain. Kondisi akhir kelelahan, waktu mulai tidur utama, dan sebagian memori slot dari hari sebelumnya diteruskan ke hari berikutnya. GA dijalankan selama 200 generasi per hari dan pencarian lokal dijalankan sesuai alur implementasi pada Bab IV; seed harian menggunakan seed dasar 42 ditambah indeks hari. Ringkasan metrik dihitung dari keluaran harian eksperimen final. Tabel 4.3 merangkum keluaran, komponen fitness, dan ukuran sebarannya.')


setp('Rata-rata fitness hasil optimasi adalah',
     'Rata-rata fitness hasil optimasi adalah −2,2101, sedangkan mediannya 6,9072. Perbedaan besar antara keduanya terjadi karena sebaran memiliki pencilan negatif: nilai minimum −1346,9761 dan maksimum 7,7997. Oleh sebab itu, rata-rata saja tidak mewakili hari yang paling umum. Nilai fitness baseline observasi juga memiliki rata-rata −158,9597 dan median −212,9600; rata-rata *fitness delta* sebesar 156,7496 perlu dibaca bersama median delta 0,9759 dan rentang −183,8144 sampai 1476,8317. Rata-rata delta yang tinggi terutama dipengaruhi beberapa baseline dengan penalti sangat besar, bukan berarti skor setiap hari meningkat sekitar 156 poin.')

p=get('Rata-rata fitness hasil optimasi adalah')
q=after(p,'Secara berpasangan, fitness optimasi lebih tinggi daripada baseline pada 399 hari (90,48%), sama pada 29 hari, dan lebih rendah pada 13 hari. Jadi, hasil agregat mendukung perbaikan skor pada mayoritas hari, tetapi tidak menunjukkan bahwa sistem selalu memperbaiki setiap jadwal. Rata-rata durasi tidur hasil optimasi 6,5045 jam per hari, sementara baseline observasi 6,5000 jam. Selisih durasi rata-rata yang kecil menunjukkan bahwa perubahan skor tidak terutama berasal dari penambahan jumlah jam tidur; fungsi tujuan juga memperhitungkan pemulihan, kepadatan aktivitas, penalti kualitas, dan regularitas.')
q2=after(q,'Rata-rata perubahan sebesar 3,1474 slot adalah jumlah slot tidur yang berbeda antara hasil optimasi dan jadwal observasi, dihitung dari gabungan slot yang ditambah atau dihilangkan. Karena satu slot setara 30 menit, besaran itu setara dengan 1,5737 jam-slot perubahan kumulatif, bukan berarti waktu mulai tidur bergeser tepat 1,57 jam. Median perubahan adalah 2 slot dan nilai maksimum 26 slot. Rata-rata *fatigue* akhir 0,6957 merupakan indeks internal model yang diwariskan antarhari; nilainya berguna untuk membandingkan keluaran di dalam eksperimen ini, tetapi tidak ditafsirkan sebagai nilai klinis.')
after(q2,'Rincian menurut hari menunjukkan bahwa median fitness optimasi relatif berdekatan, yaitu 6,5604 pada Senin; 6,9783 pada Selasa; 6,9236 pada Rabu; 6,9881 pada Kamis; 7,0150 pada Jumat; 6,8483 pada Sabtu; dan 6,9608 pada Minggu. Sebaliknya, rata-rata fitness per hari adalah 6,2715; −16,8166; 0,3628; 6,6801; −18,1428; −0,1931; dan 6,3675 secara berurutan. Perbedaan mean dan median terutama muncul karena enam keluaran memiliki fitness di bawah −100: dua pada Selasa, satu pada Rabu, dua pada Jumat, dan satu pada Sabtu. Dua kasus bahkan di bawah −500, masing-masing pada Selasa dan Jumat. Untuk konteks perubahan terhadap jadwal observasi, fitness optimasi lebih tinggi pada 58/63 hari Senin, 54/63 Selasa, 56/63 Rabu, 58/63 Kamis, 55/63 Jumat, 57/63 Sabtu, dan 61/63 Minggu; hari yang tersisa pada tiap kelompok terbagi menjadi skor sama atau lebih rendah. Dengan demikian hasil mingguan tidak seragam ant hari, tetapi mayoritas hari pada setiap hari dalam pekan tetap menunjukkan fitness yang lebih tinggi.')
after(get('Rincian menurut hari menunjukkan bahwa median fitness optimasi'), 'Rata-rata durasi tidur hasil optimasi per hari berkisar dari 6,4444 jam pada Kamis hingga 6,5397 jam pada Jumat; rata-rata fatigue akhir berkisar dari 0,5841 pada Minggu hingga 0,7734 pada Rabu. Rata-rata tujuh hari secara gabungan adalah 6,5045 jam dan 0,6957. Data baseline observasi mencatat 6,5 jam pada seluruh 441 hari, sehingga selisih rerata durasi kedua jadwal hanya 0,0045 jam per hari. Hasil ini menjelaskan bahwa kenaikan fitness tidak berasal dari peningkatan durasi rata-rata yang berarti, melainkan juga dari komponen pemulihan dan penalti di fungsi tujuan. Durasi tidur yang lebih tinggi atau fatigue internal yang lebih rendah tetap tidak membuktikan dampak kesehatan tanpa pengukuran aktual.')
after(get('Rata-rata durasi tidur hasil optimasi per hari berkisar'), 'Keluaran komponen fitness memperlihatkan recovery reward rata-rata 1,3766 (median 1,4419; rentang 0,2400–2,3996). Penalti dinamis memiliki rata-rata −2,0749, median 4,2000, dan rentang −1350 sampai 4,5500; median positif kecil menunjukkan bahwa sebagian besar hasil memperoleh komponen ini positif, sedangkan nilai ekstrem negatif menurunkan rata-rata. Penalti regularitas bernilai nol pada 438 dari 441 hari dan mencapai −400 pada tiga hari; ini adalah penalti internal yang dibatasi model, bukan nilai SRI. Terdapat satu sesi nap pada 298 hari dan tidak ada sesi nap pada 143 hari. Tidak adanya nap tidak otomatis menandakan keluaran gagal karena nap merupakan komponen opsional dan fitness menilai seluruh jadwal. Durasi tidur hasil optimasi berkisar 3–10 jam per hari; nilai minimum dan maksimum tersebut merupakan keluaran model pada kasus individual, bukan rekomendasi durasi tidur bagi pengguna.')

# Extend weekly results table with robust statistics and outcome counts.
t=doc.tables[33]
newrows=[
 ('Fitness baseline observasi (median)','−212,9600'),
 ('Fitness baseline observasi (rentang)','−1474,0858 sampai 7,5709'),
 ('Fitness delta (median)','0,9759'),
 ('Rentang fitness delta','−183,8144 sampai 1476,8317'),
 ('Hari delta positif / sama / negatif','399 / 29 / 13 dari 441 hari'),
 ('Durasi tidur optimasi (median; rentang)','6,5 jam; 3 sampai 10 jam'),
 ('Recovery reward (rata-rata / median)','1,3766 / 1,4419'),
 ('Recovery reward (rentang)','0,2400 sampai 2,3996'),
 ('Penalti dinamis (rata-rata / median)','−2,0749 / 4,2000'),
 ('Penalti dinamis (rentang)','−1350 sampai 4,5500'),
 ('Penalti regularitas (rata-rata / median)','−2,7211 / 0'),
 ('Penalti regularitas (rentang)','−400 sampai 0'),
 ('Hari dengan 1 sesi nap / tanpa nap','298 / 143 dari 441 hari'),
 ('Rentang fatigue akhir','0,0975 sampai 1,7865'),
 ('Median perubahan slot tidur','2 slot'),
 ('Rentang perubahan slot tidur','0 sampai 26 slot'),
]
for a,b in newrows:
    cells=t.add_row().cells; cells[0].text=a; cells[1].text=b

# The Table 4.2 source note was stranded alone on the next page after the long table.
for i,p in enumerate(doc.paragraphs):
    if p.text.strip().startswith('Tabel 4.3 Statistik Hasil Pengujian'):
        for following in doc.paragraphs[i+1:i+3]:
            if following.text.strip()=='Sumber: Diolah penulis':
                following._element.getparent().remove(following._element)
        break

setp('Perbandingan baseline dilakukan pada 63 baris data',
     'Perbandingan antarmetode dilakukan pada hari Rabu untuk 63 responden yang sama; seed setiap subjek merupakan seed dasar 42 ditambah indeks subjek. GA SORA dijalankan selama 150 generasi dan disempurnakan dengan pencarian lokal sebagaimana urutan implementasi pada modul baseline. Baseline *greedy* mengisi blok waktu bebas terpanjang menuju target tidur, sedangkan *fixed schedule* menempatkan pola tidur nominal pukul 22.00 selama tujuh jam. Ketiga jadwal dinilai dengan fungsi evaluasi yang sama setelah mekanisme *repair* diterapkan, sehingga angka pada Tabel 4.4 adalah metrik keluaran sesudah *repair*. Perbandingan ini merupakan uji hari Rabu, bukan pembandingan tujuh hari penuh.')

# Update table 4.3 caption and add medians and distribution counts.
# Fill the existing blank caption immediately above the baseline comparison table.
empty_captions=[p for p in doc.paragraphs if p.style.name=='Caption' and not p.text.strip()]
if len(empty_captions)!=1: raise ValueError(('empty caption count',len(empty_captions)))
empty_captions[0].add_run('Tabel 4.4 Statistik dan Hasil Berpasangan Metode pada Hari Rabu')
for p in doc.paragraphs:
    if p.style.name=='Caption' and p.text.strip().startswith('Tabel 4.1 Hasil Kalibrasi Parameter'):
        p.text=p.text.replace('Tabel 4.1','Tabel 4.2',1)
    elif p.style.name=='Caption' and p.text.strip().startswith('Tabel 4.2 Statistik Hasil Pengujian Optimasi Mingguan'):
        p.text=p.text.replace('Tabel 4.2','Tabel 4.3',1)
t=doc.tables[34]
# Round table summaries to two decimals to prevent long values from breaking
# across lines; the accompanying analysis paragraphs retain four-decimal data.
for ridx,values in {
    1:('Rata-rata fitness','6,28','−52,81','−5630,96'),
    2:('Rata-rata fatigue akhir','0,74','1,18','0,89'),
    3:('Rata-rata durasi tidur (jam)','6,52','6,90','7,43'),
}.items():
    for cell,value in zip(t.rows[ridx].cells,values): cell.text=value
rows=[
 ('Median fitness','6,59','−14,33','−13,56'),
 ('Median fatigue akhir','0,72','1,18','0,74'),
 ('Median durasi tidur (jam)','6,5','7,0','7,0'),
 ('Fitness minimum','2,63','−1347,54','−75371,12'),
 ('Fitness maksimum','6,88','6,86','6,86'),
 ('Fatigue minimum','0,28','0,26','0,31'),
 ('Fatigue maksimum','1,55','1,79','1,77'),
 ('Durasi tidur minimum (jam)','6,5','6,5','2,5'),
 ('Durasi tidur maksimum (jam)','7,0','7,0','18,0'),
 ('Fitness SORA sama','—','4 dari 63','3 dari 63'),
 ('Fitness SORA lebih rendah','—','0 dari 63','1 dari 63'),
 ('Fatigue SORA lebih rendah','—','47 dari 63','38 dari 63'),
 ('Fatigue SORA sama','—','4 dari 63','4 dari 63'),
 ('Fatigue SORA lebih tinggi','—','12 dari 63','21 dari 63'),
 ('Durasi tidur SORA lebih pendek','—','49 dari 63','42 dari 63'),
 ('Durasi tidur SORA sama','—','14 dari 63','13 dari 63'),
 ('Durasi tidur SORA lebih panjang','—','0 dari 63','8 dari 63'),
]

setp('Rata-rata fitness SORA adalah',
     'Pada 63 pasangan data hari Rabu, rata-rata fitness SORA adalah 6,2767, dibandingkan −52,8108 untuk *greedy* dan −5630,9624 untuk *fixed*. Median masing-masing adalah 6,5881; −14,3300; dan −13,5598. Nilai minimum–maksimum memperlihatkan pengaruh pencilan: SORA 2,6283–6,8822, *greedy* −1347,5356–6,8552, dan *fixed* −75371,1200–6,8552. Secara berpasangan, fitness SORA lebih tinggi daripada *greedy* pada 59 subjek, sama pada 4, dan tidak pernah lebih rendah; dibanding *fixed*, SORA lebih tinggi pada 59, sama pada 3, dan lebih rendah pada 1 subjek. SORA menjadi metode dengan fitness tertinggi terhadap kedua baseline sekaligus pada 58 dari 63 subjek. Jadi, temuan mendukung keunggulan fitness SORA pada mayoritas sampel, tetapi menyatakan kemenangan mutlak pada setiap orang akan keliru.')

setp('Hasil tersebut mendukung keunggulan SORA',
     'Rata-rata *fatigue* SORA (0,7354) lebih rendah daripada *greedy* (1,1828) dan *fixed* (0,8908); mediannya masing-masing 0,7195; 1,1800; dan 0,7415. Pada perbandingan berpasangan, fatigue SORA lebih rendah pada 47 subjek dan sama pada 4 subjek dibanding *greedy*; pada 12 subjek fatigue SORA lebih tinggi. Dibanding *fixed*, fatigue SORA lebih rendah pada 38 subjek, sama pada 4, dan lebih tinggi pada 21 subjek. Rentang fatigue adalah 0,2824–1,5500 untuk SORA, 0,2650–1,7900 untuk *greedy*, dan 0,3107–1,7700 untuk *fixed*. Arah rata-rata, median, dan mayoritas pasangan selaras: nilai fatigue SORA lebih rendah secara agregat dan pada mayoritas pasangan, meski tidak pada semua subjek.')

p=get('Rata-rata fatigue SORA')
q=after(p,'Rata-rata durasi tidur SORA 6,5159 jam (median 6,5), dibandingkan *greedy* 6,9048 jam (median 7,0) dan *fixed* 7,4286 jam (median 7,0). SORA memiliki durasi lebih pendek daripada *greedy* pada 49 subjek, sama pada 14, dan tidak lebih panjang pada pasangan mana pun. Dibanding *fixed*, durasi SORA lebih pendek pada 42 subjek, sama pada 13, dan lebih panjang pada 8. Rentang durasi masing-masing adalah 6,5–7,0; 6,5–7,0; dan 2,5–18,0 jam. Baseline *fixed* ditetapkan nominal tujuh jam sebelum *repair*, tetapi proses *repair* dapat memperluas blok tidur dan menjembatani celah pendek, sehingga keluaran aktualnya tidak selalu tujuh jam. Perbandingan durasi ini karena itu menggambarkan jadwal sesudah *repair*, bukan baseline tujuh jam yang dipertahankan tanpa perubahan.')
q2=after(q,'Ketiga jadwal dinilai dengan fungsi fitness yang sama; angka fitness adalah skor komposit tanpa satuan, bukan durasi atau tingkat kesalahan. Rata-rata *fixed* (−5630,9624) jauh di bawah mediannya (−13,5598), dan nilai minimumnya −75371,1200, sehingga beberapa pencilan mendominasi rerata. Komponen pencilan minimum dapat ditelusuri: sesudah *repair*, jadwal *fixed* tersebut berisi 36 slot tidur atau 18 jam, sedangkan target observasinya 13 slot atau 6,5 jam. Tabel 4.5 memisahkan perhitungan penalti serta komponen skor tersebut dari uraian agar setiap ekspresi matematis terbaca sebagai baris tersendiri. Skor ekstrem berasal dari keluaran *repair* yang membengkak dan bobot penalti durasi, bukan kesalahan pembacaan satuan. Perbedaan mean–median itu menjelaskan mengapa median, rentang, dan hitungan berpasangan perlu dibaca bersama mean. Setelah *repair*, baseline *fixed* dapat berubah dari target nominal tujuh jam; batas ini membatasi interpretasi keunggulan antarmetode.')

setp('Pada eksperimen mingguan, semua 63 responden',
     'Pada eksperimen mingguan, seluruh 63 responden memperoleh keluaran tujuh hari, tetapi kata “berhasil” pada status *optimized* hanya menunjukkan bahwa optimasi harian selesai menghasilkan jadwal. Rata-rata fitness −2,2101 berbeda jauh dari median 6,9072 karena beberapa kasus memperoleh penalti besar. Salah satu nilai terendah adalah −1346,9761 pada hari Jumat; komponen penalti dinamisnya −1350 dan jumlah sesi *nap* nol. Pada hari dengan jadwal wajib yang sangat padat, sistem masuk *survival mode*. Aturan implementasi memberi penalti saat tiga sesi nap tidak terbentuk; nilai komponen tersebut dituliskan terpisah pada Tabel 4.5. Jadi pencilan ini dapat ditelusuri ke aturan eksplisit fungsi fitness, bukan salah salin angka. Namun, besarnya penalti juga menunjukkan bahwa bobot aturan tersebut sangat dominan dan perlu dievaluasi ulang.')

setp('Pada perbandingan hari Rabu, SORA memiliki',
     'Pada perbandingan hari Rabu, SORA memperoleh fitness lebih tinggi daripada *greedy* pada 59 subjek dan seri pada 4; terhadap *fixed*, SORA lebih tinggi pada 59, seri pada 3, dan lebih rendah pada 1. Fitness SORA juga menjadi yang tertinggi dari ketiga metode pada 58 subjek. Untuk fatigue, SORA lebih rendah pada 47 pasangan terhadap *greedy* dan 38 terhadap *fixed*; angka tersebut tidak berarti semua pengguna memperoleh penurunan fatigue. Ketiga metrik memberi gambaran berbeda: fitness adalah fungsi tujuan komposit, fatigue adalah indeks internal model, sedangkan durasi hanya menghitung jumlah slot tidur. Hasil yang baik pada satu metrik tidak otomatis berarti nilai terbaik pada metrik lain.')

setp('Hasil ini merupakan evaluasi komputasional',
     'Hasil yang dibahas merupakan evaluasi komputasional dengan parameter eksperimen yang tercatat, bukan validasi klinis atau pengukuran langsung kualitas tidur. Penalti regularitas internal bukan perhitungan Sleep Regularity Index (SRI) standar. Sebagai contoh, delta fitness terendah −183,8144 berkaitan dengan penalti regularitas −400, yaitu batas penalti yang diterapkan ketika waktu mulai tidur berbeda jauh dari jangkar hari sebelumnya. Mekanisme ini menjelaskan penurunan skor tersebut, tetapi belum membuktikan bahwa pola tidur secara biologis tidak teratur. SORA juga tidak menghasilkan diagnosis medis atau bukti sebab-akibat terhadap produktivitas.')

p=get('Hasil yang dibahas merupakan evaluasi komputasional')
q=after(p,'Validitas generalisasi masih dibatasi oleh 63 responden, horizon pengujian tujuh hari, satu konfigurasi parameter utama, dan satu seed acak. Kalibrasi dilakukan pada lima subjek terpisah, sedangkan konfigurasi kalibrasi terbaik tidak dipakai pada eksperimen utama dan seed antarkonfigurasi kalibrasi berbeda. Perbandingan baseline hanya menguji hari Rabu dan baseline *fixed* berubah setelah *repair*. Analisis yang dilaporkan bersifat deskriptif; berkas hasil final tidak menyertakan uji inferensial atau interval kepercayaan, sehingga perbedaan angka tidak boleh disebut signifikan secara statistik. Batas-batas tersebut harus diperhitungkan saat membaca hasil dan menjadi sasaran perbaikan metodologis sebelum menyimpulkan kinerja pada populasi atau kondisi lain.')

# Chapter V: conclusions map to questions/results; do not overclaim clinical efficacy.
setp('Berdasarkan implementasi dan eksperimen pada data final',
     'Berdasarkan implementasi dan eksperimen pada data final, kesimpulan penelitian disusun untuk menjawab tujuan optimasi jadwal serta membatasi klaim pada bukti yang diukur.')
setp('SORA berhasil mengoptimasi jadwal tidur tujuh hari',
     'SORA berhasil menghasilkan jadwal tidur komputasional tujuh hari untuk 63 responden unik dari 63 baris data. Seluruh 441 evaluasi harian berstatus *optimized* dengan seed 42 dan parameter eksperimen utama populasi 50, mutasi 0,10, serta *crossover* 0,80. Keberhasilan ini menunjukkan bahwa alur aplikasi dapat memproses data dan menghasilkan keluaran untuk setiap hari; status tersebut tidak dengan sendirinya menyatakan bahwa jadwal merupakan rekomendasi klinis.')
setp('Pada optimasi mingguan, rata-rata fitness hasil',
     'Pada optimasi mingguan, rata-rata fitness hasil optimasi adalah −2,2101 dengan median 6,9072; rata-rata durasi tidur 6,5045 jam. Fitness optimasi melebihi baseline observasi pada 399 hari (90,48%), sama pada 29 hari, dan lebih rendah pada 13 hari. Median fitness delta 0,9759 lebih mewakili perubahan tipikal daripada rerata delta 156,7496 yang dipengaruhi pencilan. Temuan ini menjawab tujuan pengukuran kualitas relatif jadwal berdasarkan fungsi fitness, bukan akurasi prediksi.')
setp('Pada perbandingan 63 data hari Rabu',
     'Pada perbandingan hari Rabu untuk 63 responden, rata-rata fitness SORA 6,2767 dan median 6,5881 lebih tinggi daripada *greedy* (rata-rata −52,8108; median −14,3300) maupun *fixed* (rata-rata −5630,9624; median −13,5598). SORA memiliki fitness lebih tinggi daripada setiap baseline masing-masing pada 59 dari 63 responden, dan menjadi yang tertinggi dibanding keduanya sekaligus pada 58 responden. Kesimpulan ini berlaku untuk konfigurasi, data, hari, dan fungsi evaluasi yang diuji; tidak berarti SORA unggul pada seluruh metrik atau setiap responden.')
setp('Grid search terpisah pada lima subjek',
     'Grid search terpisah pada lima subjek menghasilkan skor rata-rata tertinggi 2,917 pada populasi 100, mutasi 0,10, dan *crossover* 0,80. Namun, metadata eksperimen utama mencatat konfigurasi fallback populasi 50, mutasi 0,10, dan *crossover* 0,80 karena pengujian utama melewati kalibrasi. Oleh sebab itu, hasil mingguan dan baseline dalam skripsi ini tidak boleh diatribusikan ke konfigurasi pemenang grid search.')

# Add an explicit evidence boundary to conclusions.
last=get('Grid search terpisah pada lima subjek')
after(last,'Secara keseluruhan, eksperimen menunjukkan bahwa SORA mampu menghasilkan dan membandingkan jadwal menggunakan fitness komposit. Penelitian ini belum menghitung SRI standar secara penuh, belum menguji dampak jadwal pada kualitas tidur nyata atau produktivitas, dan tidak memberikan diagnosis. Klaim kesimpulan dibatasi pada keluaran komputasional dari dataset dan konfigurasi yang tersedia.')

setp('Berdasarkan hasil dan keterbatasan tersebut',
     'Berdasarkan hasil dan batas penelitian yang telah dijelaskan, pengembangan selanjutnya disarankan sebagai berikut.')
setp('Perbaiki mekanisme repair dan penalti adaptif',
     'Periksa kembali bobot penalti pada kondisi *survival mode*, terutama penalti 1350 ketika sesi *nap* tidak terbentuk. Uji beberapa skenario jadwal padat dan laporkan komponen penalti secara terpisah agar perbaikan fitness tidak hanya didorong oleh perubahan satu aturan yang sangat dominan.')
setp('Ulangi eksperimen dengan beberapa random seed',
     'Ulangi eksperimen dengan beberapa seed acak dan laporkan rata-rata, median, simpangan baku, rentang, serta hasil berpasangan. Langkah ini diperlukan untuk menilai apakah hasil yang dilaporkan stabil terhadap variasi inisialisasi Genetic Algorithm.')
setp('Bandingkan SORA dengan metode optimasi tambahan',
     'Bandingkan SORA dengan metode optimasi tambahan menggunakan data, kendala, fungsi fitness, seed, dan anggaran komputasi yang sama. Baseline *fixed* perlu dievaluasi konsistensinya: jika maksud baseline adalah tepat tujuh jam mulai pukul 22.00, tetapkan aturan agar proses *repair* tidak mengubah durasi tersebut, atau nyatakan secara eksplisit perubahan setelah *repair*.')
setp('Validasi sistem pada data pengguna nyata',
     'Perluas validasi dengan data longitudinal yang lebih beragam dan, bila tersedia, data *wearable* yang telah mendapat persetujuan etik dan persetujuan responden. Data aktual tersebut dapat digunakan untuk menguji kesesuaian prediksi jadwal dengan perilaku tidur, bukan hanya skor simulasi.')
# Add two focused development recommendations.
p=get('Perluas validasi dengan data longitudinal')
q=after(p,'Jika regularitas tidur menjadi sasaran evaluasi, hitung SRI sesuai definisi dan kebutuhan data aslinya, lalu laporkan terpisah dari penalti regularitas internal SORA. Hal ini mencegah dua ukuran yang berbeda diberi nama atau makna yang sama.')
q2=after(q,'Laporkan artefak eksperimen (metadata parameter, seed, data yang telah dianonimkan, ringkasan metrik, dan skrip yang diperlukan) secara konsisten agar pembaca dapat menelusuri asal setiap angka dan mengulang analisis tanpa membuka identitas responden.')

# Add median and paired-outcome rows to the baseline table after the existing rows.
t=doc.tables[34]
# Existing table already includes mean metrics and paired fitness counts; append robust medians and fatigue counts.
for label,sora,greedy,fixed in rows:
    cells=t.add_row().cells
    for c,value in zip(cells,(label,sora,greedy,fixed)): c.text=value
    for c,value in zip(cells,(label,sora,greedy,fixed)): c.text=value

# Update table captions and clear accidental duplicate source insertion if any.
# Ensure table header rows repeat when spanning pages.
for table in [doc.tables[32],doc.tables[33],doc.tables[34]]:
    trPr=table.rows[0]._tr.get_or_add_trPr()
    from docx.oxml import OxmlElement
    if trPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader') is None:
        trPr.append(OxmlElement('w:tblHeader'))

# The imported template's Heading 2 style adds an automatic list number on top
# of the subsection numbers already present in the manuscript text.
heading2=doc.styles['Heading 2']
if heading2._element.pPr is not None and heading2._element.pPr.numPr is not None:
    heading2._element.pPr.remove(heading2._element.pPr.numPr)
heading3=doc.styles['Heading 3']
if heading3._element.pPr is not None and heading3._element.pPr.numPr is not None:
    heading3._element.pPr.remove(heading3._element.pPr.numPr)

# The original manuscript relied on template numbering for these headings.
# The imported template's numbering was malformed (it rendered as 7.x), so
# preserve the intended chapter numbering explicitly in the heading text.
manual_numbers={
    'Latar Belakang Masalah':'1.1 ', 'Rumusan Masalah':'1.2 ',
    'Batasan Masalah':'1.3 ', 'Tujuan Penelitian':'1.4 ',
    'Manfaat Penelitian':'1.5 ', 'Tahapan Penelitian':'1.6 ',
    'Sistematika Penulisan':'1.7 ',
    'Jenis dan Rancangan Penelitian':'3.1 ', 'Data Penelitian':'3.2 ',
    'Alur Penelitian':'3.3 ', 'Pra-Pemrosesan Data':'3.4 ',
    'Prosedur Eksekusi Genetic Algorithm dan Hill Climbing':'3.5 ',
    'Optimasi Jadwal Tidur Mingguan':'3.6 ',
    'Konfigurasi Parameter Eksperimen Utama':'3.7 ',
    'Baseline Comparison':'3.8 ', 'Evaluasi Sistem':'3.9 ',
    'Implementasi Program':'3.10 ',
    'Optimasi Hari Pertama':'3.6.1 ', 'Carry-Over Kondisi Antarhari':'3.6.2 ',
    'Optimasi Hari Kedua hingga Hari Ketujuh':'3.6.3 ',
    'Penyusunan Jadwal Mingguan':'3.6.4 ', 'Hasil Optimasi Mingguan':'3.6.5 ',
    'Weekly Optimization Test':'3.9.1 ',
    '3.8 Baseline Comparison':'3.9.2 Baseline Comparison',
    'Multi-Subject Test':'3.9.3 ', 'Indikator Evaluasi':'3.9.4 ',
    'Analisis Hasil Pengujian':'3.9.5 ',
}
for p in doc.paragraphs:
    if p.text.strip() in manual_numbers:
        old=p.text.strip()
        p.text=manual_numbers[old] if old=='3.8 Baseline Comparison' else manual_numbers[old]+old

# Start each main chapter on a new page.
for p in doc.paragraphs:
    if p.style.name=='Heading 1' and p.text.strip() in {'BAB I','BAB II','BAB III','BAB IV','BAB V'}:
        p.paragraph_format.page_break_before=True

# Remove the empty end-of-methodology heading; its section had no body text.
for p in list(doc.paragraphs):
    if p.style.name=='Heading 2' and p.text.strip()=='Ringkasan Metodologi':
        p._element.getparent().remove(p._element)

# Give the metric labels more room so the final rows do not spill alone to a page.
from docx.shared import Cm
doc.tables[33].columns[0].width=Cm(8.5)
doc.tables[33].columns[1].width=Cm(7.0)
doc.tables[34].columns[0].width=Cm(6.0)
for idx in (1,2,3): doc.tables[34].columns[idx].width=Cm(3.15)
from docx.shared import Pt
for row in doc.tables[34].rows:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs: run.font.size=Pt(9)

# Chapter IV must report the realized product before presenting test results.
# Keep the implementation description within one existing chapter subsection;
# do not create extra sub-subsections for mathematical models.
calibration_heading=get('4.1 Hasil Kalibrasi Parameter')
implementation_heading=before(calibration_heading,'4.1 Implementasi Sistem','Heading 2')
implementation_paragraphs=[
    'Implementasi penelitian menghasilkan SORA Academic, yaitu aplikasi optimasi jadwal tidur yang terdiri atas mesin optimasi berbasis Python, layanan REST API menggunakan FastAPI, dan antarmuka web React. Mesin optimasi merealisasikan rancangan pada Bab III menjadi alur yang dapat menerima data aktivitas, membentuk jadwal harian, menjalankan optimasi, dan mengembalikan jadwal beserta metriknya. Antarmuka dan API menghubungkan fungsi tersebut dengan pengguna. Implementasi ini merupakan prototipe penelitian yang dijalankan pada lingkungan lokal; hasilnya tidak menyatakan bahwa sistem telah dipasang sebagai layanan produksi atau digunakan untuk memberi rekomendasi klinis.',
    'Parser membaca dataset kuesioner final kuisoner (Jawaban) - Form Responses 1.csv. Kolom nama sumber dan kolom aktivitas untuk tujuh hari dipetakan menjadi Nama, Kerja_<Hari>, Kuliah_<Hari>, dan Tidur_<Hari>. Nilai kosong, nan, null, none, atau tanda hubung untuk rentang aktivitas menghasilkan daftar slot kosong. Rentang waktu yang dipisahkan koma dibaca satu per satu, divalidasi dengan pola jam-menit, lalu digabungkan sebagai daftar indeks slot unik yang diurutkan. Parser juga menangani rentang yang melewati tengah malam dengan melanjutkan indeks ke awal hari berikutnya.',
    'Contoh input anonim dari format kuesioner: 07:30 - 12:30. Keluaran parser adalah indeks slot 15 sampai 24, yaitu sepuluh slot aktivitas yang masing-masing berdurasi 30 menit. Pada struktur data normalisasi, contoh tersebut disimpan pada kolom Kerja_<Hari> sebagai teks daftar indeks; identitas responden tidak diperlukan dalam contoh ini. Jika format waktu tidak cocok atau indeks slot terduplikasi/tidak berada pada rentang yang diizinkan, proses menghasilkan kesalahan validasi alih-alih meneruskan nilai yang tidak sah.',
    'Satu kromosom harian memuat 48 gen, satu gen untuk setiap slot 30 menit. Enumerasi Allele pada representation.py menetapkan SLEEP=0, WORK=1, NAP=2, REST=3, dan CLASS=4. Contoh representasi: jika aktivitas kerja berlangsung pada slot 15–24, maka gen pada indeks tersebut bernilai WORK (1); nilai pada indeks lain mengikuti aktivitas dalam kandidat jadwal. Fungsi validate_chromosome memeriksa panjang kromosom dan menolak alel di luar lima nilai tersebut. Jadwal kerja dan kuliah membentuk mandatory_core; satu slot di sekitar setiap slot wajib yang tidak bersebelahan ditambahkan sebagai transit slot. Gabungan keduanya menjadi mandatory_absolute yang tidak boleh diubah operator optimasi.',
    'Pemeriksaan constraint juga memvalidasi bahwa daftar slot merupakan bilangan bulat dalam rentang 0–47, tidak memiliki duplikasi, dan slot kerja tidak tumpang tindih dengan slot kuliah. Hari dengan lebih dari 36 slot wajib tidak dijalankan oleh optimasi mingguan dan diberi status skipped_infeasible. Jika ruang kosong hari tersebut tidak memiliki rentang kontinu sedikitnya 12 slot, konteks ditandai survival_mode; aturan fitness dan repair menggunakan perlakuan khusus untuk kondisi jadwal padat ini.',
    'Fungsi calculate_fitness pada fitness.py mengembalikan fitness, fatigue akhir, durasi tidur, awal blok tidur utama, recovery reward, dynamic penalty, regularity penalty, dan jumlah sesi nap. Recovery dihitung sepanjang slot berdasarkan aktivitas dan bobot waktu sirkadian; kerja, kuliah, dan istirahat menambah fatigue dengan laju berbeda, sedangkan tidur dan nap yang memenuhi aturan menguranginya serta dapat memberi reward. Penalti dinamis memeriksa durasi tidur terhadap minimum, maksimum, dan target; kesesuaian dengan durasi observasi; tumpang tindih slot tidur; jumlah blok; jumlah serta waktu nap. Penalti regularitas membandingkan awal tidur utama dengan anchor hari sebelumnya dan dibatasi maksimum 400. Skor akhir mengikuti ekspresi pada Tabel 4.1; fitness adalah skor internal fungsi tujuan, bukan akurasi maupun ukuran klinis.',
    'Pada setiap generasi Genetic Algorithm, populasi dirangking dari fitness tertinggi. Individu elit dipertahankan, sedangkan indeks dua induk dipilih dari kelompok teratas. Dengan probabilitas crossover, satu anak dibentuk melalui single-point crossover pada titik potong yang dipilih di antara slot 10 dan 37; jika crossover tidak terjadi, kromosom induk pertama disalin. Mutasi hanya menyentuh slot nonwajib dan memakai distribusi alel yang berbeda untuk jendela nap dan slot di luar jendela nap. Anak kemudian melalui repair sebelum masuk ke generasi berikutnya. Pengujian utama menggunakan populasi 50, mutasi 0,10, crossover 0,80, 200 generasi, dan seed 42 ditambah indeks hari.',
    'Hill Climbing membentuk kandidat tetangga dari pertukaran berurutan antara alel tidur, nap, atau istirahat yang berbeda serta perubahan alel pada slot nonwajib. Setiap kandidat diperbaiki dan dinilai kembali; kandidat diterima hanya jika fitness-nya lebih tinggi daripada skor terbaik sementara. Iterasi berhenti setelah batas iterasi tercapai atau satu putaran penuh tidak menemukan perbaikan. Pada implementasi aktual, optimizer menjalankan Hill Climbing pada individu terbaik setelah GA; weekly.py kemudian memanggil run_hill_climbing sekali lagi pada hasil tersebut. Karena itu proses mingguan memiliki dua pemanggilan pencarian lokal, masing-masing dengan batas konfigurasi 20 iterasi, bukan satu pemanggilan total 20 iterasi. Perilaku ini dituliskan apa adanya agar uraian dan cuplikan kode konsisten dengan eksperimen yang menghasilkan data final.',
    'Fungsi run_weekly_optimization_with_history memproses Senin sampai Minggu secara berurutan. Untuk tiap hari fungsi membaca slot kerja, kuliah, dan tidur observasi, membentuk optimizer dengan fatigue awal, anchor awal tidur, debt installment, memori 16 slot terakhir, dan seed harian, lalu mengecek feasibility. Jika layak, proses menjalankan GA dan pencarian lokal, menghitung metrik, serta menyimpan jadwal hasil dan jadwal observasi. Nilai fatigue akhir, awal blok tidur, dan 16 gen terakhir diteruskan ke hari berikutnya. Hasil harian memuat fitness, fitness observasi, fitness_delta, sleep_hours, sleep_change_slots, recovery, penalti, regularity_penalty, nap_sessions, dan status.',
    'Endpoint POST /api/weekly/optimize di api.py menerima nama subjek, dataset, parameter, dan seed melalui WeeklyRequest. Endpoint meneruskan permintaan ke fungsi optimasi mingguan dan mengembalikan target_name, schedule, observed_schedule, metrics, serta history. Jika subjek tidak ditemukan, endpoint mengirim respons HTTP 404. Contoh keluaran secara struktural berisi tujuh jadwal yang masing-masing mempunyai 48 alel, tujuh objek metrik harian, dan riwayat generasi; nilai contoh dalam naskah tidak menggunakan nama atau identitas responden.',
    'Antarmuka React menyajikan halaman Beranda, Dataset, Mingguan, Baseline, dan Kalibrasi. Pada halaman Mingguan, pengguna memilih subjek dan menjalankan optimasi; hasilnya ditampilkan sebagai perbandingan jadwal observasi dan jadwal optimasi dalam garis waktu, disertai perubahan slot, durasi tidur, nap, fitness, serta riwayat generasi. Halaman Baseline meminta API membandingkan metode SORA, greedy, dan fixed untuk subjek terpilih. Halaman Kalibrasi menjalankan endpoint grid search dan menerima ringkasan parameter terbaik. Proyek juga menyediakan antarmuka alternatif Streamlit pada frontend/app.py dengan menu Dataset, Weekly, Baseline, Calibration, dan Multi-Subject. Angka pada penelitian tetap merujuk pada artefak pipeline final yang parameternya tercatat.',
    'Peta kode inti mencakup parser.py untuk pembacaan dan normalisasi; representation.py dan constraints.py untuk alel, kromosom, serta kendala; fitness.py untuk penilaian; population.py, selection.py, crossover.py, mutation.py, repair.py, dan optimizer.py untuk Genetic Algorithm; local_search.py untuk Hill Climbing; weekly.py untuk alur tujuh hari; serta api.py untuk endpoint. Cuplikan kode inti pada bagian berikut diambil langsung dari berkas tersebut. Kode lengkap tetap berada di direktori proyek sora_academic; listing dipilih seperlunya agar tiap tahap dapat ditelusuri tanpa memenuhi badan skripsi dengan seluruh kode sumber.',
    'Hasil implementasi pada data final mencakup pemrosesan 63 responden selama tujuh hari, 441 keluaran harian berstatus optimized, dan 63 pasangan keluaran hari Rabu untuk baseline. Status optimized menunjukkan bahwa proses komputasi menghasilkan keluaran, bukan jaminan jadwal ideal atau kelayakan klinis.'
]
anchor=implementation_heading
for idx,text in enumerate(implementation_paragraphs):
    headings={1:'4.1.1 Implementasi Dataset dan Normalisasi',3:'4.1.2 Representasi Kromosom dan Constraint',5:'4.1.3 Implementasi Fungsi Fitness',6:'4.1.4 Implementasi Genetic Algorithm',7:'4.1.5 Implementasi Hill Climbing',8:'4.1.6 Implementasi Optimasi Mingguan dan API',10:'4.1.7 Implementasi Antarmuka dan Keluaran'}
    if idx in headings: anchor=after(anchor,headings[idx],'Heading 3')
    anchor=after(anchor,text)

# Aggregate visuals based solely on anonymized final results.
from PIL import Image, ImageDraw, ImageFont
figdir=Path(r'D:\Baretta\Skripsi\program\sora_academic\thesis_dump\figures')
figdir.mkdir(exist_ok=True)
def ff(size,bold=False):
    path=r'C:\Windows\Fonts\arialbd.ttf' if bold else r'C:\Windows\Fonts\arial.ttf'
    try: return ImageFont.truetype(path,size)
    except OSError: return ImageFont.load_default()
flow=Image.new('RGB',(1500,650),'white'); dr=ImageDraw.Draw(flow)
dr.text((60,30),'Alur implementasi SORA Academic',font=ff(36,True),fill='#17324D')
steps=[('Data kuesioner dan pemilihan subjek',60,130),('Parser dan normalisasi menjadi 48 slot/hari',420,130),('Representasi jadwal dan kendala wajib',780,130),('Genetic Algorithm',1140,130),('Repair dan Hill Climbing',940,390),('Evaluasi dan keluaran jadwal serta metrik',500,390),('Carry-over kondisi Senin–Minggu',60,390)]
for label,x,y in steps:
    dr.rounded_rectangle((x,y,x+300,y+135),radius=16,fill='#EAF1F8',outline='#284B63',width=3)
    words=label.split(); lines=[]; line=''
    for word in words:
        test=(line+' '+word).strip()
        if dr.textbbox((0,0),test,font=ff(22))[2] > 260: lines.append(line); line=word
        else: line=test
    if line: lines.append(line)
    yy=y+35
    for item in lines:
        tw=dr.textbbox((0,0),item,font=ff(22))[2]; dr.text((x+(300-tw)//2,yy),item,font=ff(22),fill='#172B3A'); yy+=30
for x in (360,720,1080):
    dr.line((x+5,198,x+48,198),fill='#284B63',width=4); dr.polygon([(x+48,198),(x+34,189),(x+34,207)],fill='#284B63')
dr.line((1290,265,1290,320,1090,320,1090,385),fill='#284B63',width=4)
dr.polygon([(1090,385),(1081,370),(1099,370)],fill='#284B63')
dr.line((940,458,800,458),fill='#284B63',width=4); dr.polygon([(800,458),(815,449),(815,467)],fill='#284B63')
dr.line((500,458,360,458),fill='#284B63',width=4); dr.polygon([(360,458),(375,449),(375,467)],fill='#284B63')
flow_path=figdir/'gambar_4_1_alur_sora.png'; flow.save(flow_path)

days=['Senin','Selasa','Rabu','Kamis','Jumat','Sabtu','Minggu']
means=[6.2715,-16.8166,.3628,6.6801,-18.1428,-.1931,6.3675]
medians=[6.5604,6.9783,6.9236,6.9881,7.0150,6.8483,6.9608]
wins=[58,54,56,58,55,57,61]; ties=[2,6,7,3,7,2,2]; losses=[3,3,0,2,1,4,0]
chart=Image.new('RGB',(1500,820),'white'); dr=ImageDraw.Draw(chart)
dr.text((55,25),'Fitness harian SORA dan hasil berpasangan terhadap observasi',font=ff(32,True),fill='#17324D')
dr.text((75,95),'A. Rerata dan median fitness (titik merah = median)',font=ff(23,True),fill='#263746')
left,top,width,height=150,145,1260,280; zero=top+height//2
for v in [-20,-10,0,10]:
    y=zero-int(v*7.5); dr.line((left,y,left+width,y),fill='#D5DCE3',width=1); dr.text((85,y-10),str(v),font=ff(16),fill='#445')
for i,(day,m,med) in enumerate(zip(days,means,medians)):
    x=left+70+i*175; by=zero-int(m*7.5)
    dr.rectangle((x,min(zero,by),x+58,max(zero,by)),fill='#5286A6')
    my=zero-int(med*7.5); dr.ellipse((x+20,my-8,x+36,my+8),fill='#D1495B')
    dr.text((x-10,top+height+9),day,font=ff(16),fill='#263746')
dr.text((75,480),'B. Pasangan fitness: optimasi lebih tinggi / sama / lebih rendah (n=63)',font=ff(23,True),fill='#263746')
for i,(day,w,t,l) in enumerate(zip(days,wins,ties,losses)):
    y=530+i*36; x=180; dr.text((75,y),day,font=ff(16),fill='#263746')
    for n,color in ((w,'#39835D'),(t,'#D7A839'),(l,'#C65D55')):
        dr.rectangle((x,y,x+n*12,y+24),fill=color); x+=n*12
    dr.text((970,y),f'{w} / {t} / {l}',font=ff(16),fill='#263746')
chart_path=figdir/'gambar_4_2_fitness_harian.png'; chart.save(chart_path)
def insert_picture_after(paragraph,path,width=6.0):
    q=after(paragraph,''); q.alignment=1
    q.add_run().add_picture(str(path),width=Inches(width)); return q
chart_anchor=insert_picture_after(get('Rata-rata durasi tidur hasil optimasi per hari berkisar'),chart_path)
chart_anchor.paragraph_format.keep_with_next=True
chart_anchor=after(chart_anchor,'Gambar 4.2 Rerata, median, dan hasil berpasangan fitness menurut hari','Caption')
chart_anchor.paragraph_format.keep_with_next=True
chart_anchor=after(chart_anchor,'Sumber: Diolah penulis dari 441 keluaran eksperimen mingguan','Normal')

# Present representative source code from the running implementation in Bab IV.
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
def source_listing(after_paragraph,label,lines):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    label_p=after(after_paragraph,label)
    label_p.runs[0].bold=True
    label_p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    label_p.paragraph_format.keep_with_next=True
    code_p=after(label_p,'')
    code_p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    code_p.paragraph_format.keep_together=True
    code_p.paragraph_format.left_indent=Cm(0.2)
    code_p.paragraph_format.first_line_indent=Cm(0)
    code_p.paragraph_format.space_before=Pt(0)
    code_p.paragraph_format.space_after=Pt(6)
    code_p.paragraph_format.line_spacing=1.0
    for line_no,line in enumerate(lines):
        run=code_p.add_run(line)
        run.font.name='Consolas'; run.font.size=Pt(7.5)
        if line_no < len(lines)-1: run.add_break()
    return code_p

def insert_picture_after(paragraph,path,width=6.0):
    q=after(paragraph,''); q.alignment=1
    q.add_run().add_picture(str(path),width=Inches(width)); return q
anchor=insert_picture_after(anchor,flow_path)
anchor.paragraph_format.page_break_before=True
anchor.paragraph_format.keep_with_next=True
anchor=after(anchor,'Gambar 4.1 Alur implementasi SORA Academic','Caption')
anchor.paragraph_format.keep_with_next=True
anchor=after(anchor,'Sumber: Diolah penulis berdasarkan implementasi SORA Academic','Normal')
anchor=after(anchor,'4.1.8 Pseudocode dan Cuplikan Kode Inti','Heading 3')
anchor=source_listing(anchor,'Pseudocode 4.1 Alur pemrosesan satu subjek selama tujuh hari (pseudocode, bukan kode sumber)',[
    'BACA berkas kuesioner dan normalisasikan kolom aktivitas',
    'PILIH satu baris subjek; inisialisasi fatigue, anchor, dan memori slot',
    'UNTUK setiap hari Senin sampai Minggu:',
    '  PARSE slot kerja, kuliah, dan tidur observasi',
    '  VALIDASI irisan kerja-kuliah dan kelayakan slot wajib',
    '  BENTUK konteks optimasi dengan seed dasar + indeks hari',
    '  BANGUN jadwal observasi sebagai pembanding',
    '  JIKA layak: jalankan GA selama 200 generasi dan lakukan repair',
    '  JALANKAN Hill Climbing internal optimizer pada kandidat terbaik',
    '  JALANKAN pemanggilan Hill Climbing eksplisit pada weekly.py',
    '  HITUNG fitness, fatigue, recovery, penalti, dan perubahan slot',
    '  SIMPAN jadwal; teruskan fatigue, anchor, serta 16 slot terakhir',
 'KEMBALIKAN jadwal mingguan, metrik harian, dan riwayat konvergensi',
])
anchor=after(anchor,'Cuplikan kode berikut disalin langsung dari berkas sumber yang disebutkan pada judul listing. Setiap cuplikan menunjukkan operasi aktual; uraian algoritmik di atas adalah pseudocode untuk merangkum urutan proses.','Normal')
def exact_source(relative_path,start_marker,end_marker,include_end=False):
    lines=(Path(r'D:\Baretta\Skripsi\program\sora_academic')/relative_path).read_text(encoding='utf-8').splitlines()
    starts=[i for i,line in enumerate(lines) if line.startswith(start_marker)]
    if len(starts)!=1:
        raise ValueError((relative_path,start_marker,starts))
    end=next((i for i in range(starts[0]+1,len(lines)) if lines[i].startswith(end_marker)),None)
    if end is None:
        raise ValueError((relative_path,start_marker,end_marker))
    return lines[starts[0]:end+(1 if include_end else 0)]
anchor=source_listing(anchor,'Cuplikan kode 4.1 Parser rentang waktu (src/sora/data/parser.py)',exact_source('src/sora/data/parser.py','def parse_time_ranges','def normalize_questionnaire_dataframe'))
anchor=source_listing(anchor,'Cuplikan kode 4.2 Alel aktivitas dalam kromosom (src/sora/genetic_algorithm/representation.py)',exact_source('src/sora/genetic_algorithm/representation.py','class Allele','VALID_ALLELES ='))
anchor=source_listing(anchor,'Cuplikan kode 4.3 Perhitungan skor fitness akhir (src/sora/genetic_algorithm/fitness.py)',exact_source('src/sora/genetic_algorithm/fitness.py','    fitness = (context.sleep_weight','    return FitnessResult('))
anchor=source_listing(anchor,'Cuplikan kode 4.4 Siklus generasi dan operator GA (src/sora/genetic_algorithm/optimizer.py)',exact_source('src/sora/genetic_algorithm/optimizer.py','        for generation_idx in range(generations):','        if population:'))
anchor=source_listing(anchor,'Cuplikan kode 4.5 Urutan repair kromosom (src/sora/genetic_algorithm/repair.py)',exact_source('src/sora/genetic_algorithm/repair.py','def repair_chromosome','    return chromosome',include_end=True))
anchor=source_listing(anchor,'Cuplikan kode 4.6 Pencarian lokal Hill Climbing (src/sora/hill_climbing/local_search.py)',exact_source('src/sora/hill_climbing/local_search.py','def run_hill_climbing','    return best_chromosome',include_end=True))
anchor=source_listing(anchor,'Cuplikan kode 4.7 Pemanggilan GA dan Hill Climbing dari weekly.py (src/sora/experiments/weekly.py)',exact_source('src/sora/experiments/weekly.py','        best_ga, day_history = optimizer.run_genetic_algorithm_with_history(','        result = optimizer.calculate_fitness(best_schedule)',include_end=True))
anchor=source_listing(anchor,'Cuplikan kode 4.8 Endpoint optimasi mingguan (api.py)',exact_source('api.py','@app.post("/api/weekly/optimize")','@app.post("/api/baseline/compare")'))

# The implemented fitness expression is shown as a one-column equation table.
fitness_anchor=get('Fungsi calculate_fitness pada fitness.py')
fitness_caption=after(fitness_anchor,'Tabel 4.1 Persamaan Fitness yang Diterapkan dalam Implementasi','Caption')
fitness_table=doc.add_table(rows=1,cols=1); fitness_table.style='Table Grid'
fitness_cell=fitness_table.cell(0,0); fitness_cell.text='F = w_s × (R + 0,2D) + P_d + P_r'
for p in fitness_cell.paragraphs:
    p.alignment=1
    for run in p.runs: run.font.name='Cambria Math'; run.font.size=Pt(12)
fitness_caption._p.addnext(fitness_table._tbl)
fitness_table_note_el=OxmlElement('w:p'); fitness_table._tbl.addnext(fitness_table_note_el)
fitness_table_note=Paragraph(fitness_table_note_el,fitness_table._parent); fitness_table_note.style='Normal'
fitness_table_note.add_run('Keterangan: w_s adalah bobot tidur; R recovery reward; D jumlah gen tidur (slot); P_d dynamic penalty; dan P_r regularity penalty. Persamaan ini setara dengan ekspresi yang dieksekusi pada fitness.py.').italic=True

# Mathematical expressions are isolated in a one-column table, separate from
# the explanatory prose as requested by the author.
formula_anchor=get('Ketiga jadwal dinilai dengan fungsi fitness yang sama;')
formula_caption=after(formula_anchor,'Tabel 4.5 Rincian Perhitungan Komponen Fitness pada Kasus Ekstrem','Caption')
formula_table=doc.add_table(rows=0,cols=1)
formula_table.style='Table Grid'
for expression in [
    '−200 × (36 − 18)² = −64.800',
    '−20 × (36 − 13)² = −10.580',
    '12 × 0,35 = 4,20',
    '−64.800 − 10.580 + 4,20 = −75.375,80',
    '−75.375,80 + 4,68 = −75.371,12',
    '−150 × 3² = −1.350',
]:
    cell=formula_table.add_row().cells[0]
    cell.text=expression
    for p in cell.paragraphs:
        p.alignment=1
        for run in p.runs: run.font.name='Cambria Math'; run.font.size=Pt(11)
formula_caption._p.addnext(formula_table._tbl)
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
note_el=OxmlElement('w:p'); formula_table._tbl.addnext(note_el)
formula_note=Paragraph(note_el,formula_table._parent); formula_note.style='Normal'
formula_note.add_run('Keterangan: baris pertama sampai kelima menjabarkan skor minimum metode fixed pada hari Rabu; baris terakhir menunjukkan penalti survival mode pada keluaran mingguan hari Jumat.').italic=True

# The Markdown source explicitly separates the experiment pipeline demo from
# the baseline comparison. Insert the reproducibility artifacts after weekly
# outcomes and before baseline analysis.
baseline_heading=get('4.3 Hasil Pengujian Baseline Comparison')
pipeline_heading=before(baseline_heading,'4.4 Hasil Demo Pipeline Skripsi','Heading 2')
pipeline_texts=[
    'Pipeline demo skripsi dijalankan untuk mengemas keluaran eksperimen final menjadi artefak yang dapat diperiksa dan digunakan dalam penyusunan laporan. Berkas README.md menjelaskan isi paket, INSTRUCTIONS.md memuat petunjuk penggunaan, METADATA.json mencatat parameter dan seed, sedangkan REPORT.md dan REPORT.csv menyajikan ringkasan hasil. summary_report.csv merangkum keluaran mingguan; direktori results memuat ringkasan kalibrasi, perbandingan multi-subjek, metrik harian, dan jadwal mingguan. Artefak per responden mengandung data sensitif sehingga tidak ditampilkan dengan identitas pada naskah.',
    'Metadata eksperimen utama mencatat seed 42, ukuran populasi 50, laju mutasi 0,10, laju crossover 0,80, serta sumber parameter default_fallback karena eksekusi utama menggunakan --skip-calibration. Ringkasan final memuat 63 responden dan 441 baris keluaran harian. Keterlacakan ini memungkinkan pembaca memeriksa konsistensi antara parameter, hasil agregat, dan berkas eksperimen tanpa menyamakan pipeline demo dengan kalibrasi parameter.'
]
pipeline_anchor=pipeline_heading
for text in pipeline_texts: pipeline_anchor=after(pipeline_anchor,text)

# Shift section numbers to match the supplied Bab IV outline.
for old,new in [
    ('4.1 Hasil Kalibrasi Parameter','4.2 Hasil Kalibrasi Parameter'),
    ('4.2 Hasil Pengujian Optimasi Mingguan','4.3 Hasil Pengujian Optimasi Mingguan'),
    ('4.3 Hasil Pengujian Baseline Comparison','4.5 Perbandingan dengan Baseline'),
    ('4.4 Pembahasan','4.6 Pembahasan Hasil'),
]:
    matches=[p for p in doc.paragraphs if p.style.name=='Heading 2' and p.text.strip()==old]
    if len(matches)!=1: raise ValueError(('Chapter IV heading',old,len(matches)))
    matches[0].text=new

# Add the subparts present in the author's Markdown base.
babv_heading=next(p for p in doc.paragraphs if p.style.name=='Heading 1' and p.text.strip()=='BAB V')
summary_heading=before(babv_heading,'4.7 Ringkasan Pembahasan','Heading 2')
summary_text=after(summary_heading,'Secara umum, hasil final menunjukkan bahwa SORA dapat memproses data kuesioner, menghasilkan jadwal untuk tujuh hari, membandingkan keluaran dengan baseline, dan mengemas ringkasan eksperimen. Pada 399 dari 441 hari fitness hasil optimasi lebih tinggi daripada jadwal observasi, sedangkan pada perbandingan Rabu fitness SORA tertinggi dibanding kedua baseline secara simultan pada 58 dari 63 responden. Temuan ini merupakan perbandingan skor komputasional dalam konfigurasi penelitian; temuan tersebut tidak membuktikan peningkatan kualitas tidur aktual atau dampak klinis.')
limit_heading=after(summary_text,'4.8 Keterbatasan Pengujian','Heading 2')
limit_text=after(limit_heading,'Hasil pengujian dibatasi pada 63 responden dan horizon tujuh hari, dengan satu konfigurasi utama serta satu seed acak. Analisis yang tersedia bersifat deskriptif dan tidak menyertakan uji inferensial atau interval kepercayaan. Perbandingan antarmetode hanya dilakukan pada hari Rabu; selain itu, proses repair dapat mengubah jadwal fixed nominal tujuh jam. Implementasi menggunakan regularity penalty internal dan belum menghitung SRI standar, serta belum memvalidasi jadwal terhadap pengukuran tidur aktual. Oleh karena itu, fitness dan fatigue dibahas sebagai keluaran internal sistem, bukan ukuran klinis atau bukti manfaat kesehatan.')

# Bab V follows the supplied structure: conclusion, limitations, recommendations,
# and a concise closing statement.
suggestion_heading=get('5.2 Saran')
suggestion_heading.text='5.3 Saran'
thesis_limit=before(suggestion_heading,'5.2 Keterbatasan Penelitian','Heading 2')
thesis_limit_p=after(thesis_limit,'Penelitian ini menggunakan 63 responden unik dan membatasi optimasi pada tujuh hari. Eksperimen utama menggunakan satu seed dan parameter default/fallback karena kalibrasi dilewati; baseline dibandingkan hanya pada hari Rabu. Tidak tersedia uji inferensial atau interval kepercayaan, pengukuran langsung atas kualitas tidur, maupun perhitungan SRI standar. Beberapa metrik seperti fitness, fatigue, dan regularity penalty merupakan ukuran internal fungsi evaluasi, sehingga hasil tidak dapat digeneralisasi sebagai rekomendasi medis.','Normal')
last_suggestion=next(p for p in reversed(doc.paragraphs) if p.text.strip().startswith('Laporkan artefak eksperimen'))
closing_heading=after(last_suggestion,'5.4 Penutup','Heading 2')
after(closing_heading,'Penelitian ini menghasilkan prototipe optimasi jadwal tidur dan artefak eksperimen yang dapat ditelusuri. Kesimpulan dibatasi pada data dan konfigurasi yang diuji; validasi lebih luas dan pengukuran aktual tetap diperlukan sebelum hasil digunakan untuk menyatakan manfaat terhadap tidur atau kesehatan.','Normal')

case_heading=before(get('Pada eksperimen mingguan, seluruh 63 responden'),'4.6.1 Analisis Ringkas Per Subjek','Heading 3')
after(case_heading,'Sebagai ilustrasi anonim, satu kasus dengan rerata fitness mingguan tepat pada median responden dipilih untuk ditelusuri. Pada hari Rabu, fitness SORA sebesar 7,2948, sedangkan jadwal observasi memperoleh 7,0043; selisihnya 0,2906. Keduanya memiliki durasi tidur utama 6,5 jam. Jadwal hasil optimasi menambahkan tiga slot nap pada rentang tengah hari, sekitar pukul 12.00–13.30. Contoh ini memperlihatkan bahwa skor dapat berubah tanpa perubahan durasi tidur utama, tetapi satu kasus tidak mewakili seluruh responden dan tidak membuktikan dampak tidur aktual.')

doc.core_properties.title='Optimasi Jadwal Kerja dan Istirahat Menggunakan Algoritma Genetika Pendekatan pada Pekerja dengan Pola Tidak Teratur'
doc.save(out)
print(out)

