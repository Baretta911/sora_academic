from docx import Document
p=r'D:\Baretta\Skripsi\program\sora_academic\thesis_dump\Skripsi_123220204_BarettaP_final_bab4_bab5.docx';d=Document(p)
for x in d.paragraphs:
 if x.text.strip().startswith(('BAB IV','4.','BAB V','5.')): print(x.style.name,x.text[:150])
print('TABLES',len(d.tables),'table33 rows',len(d.tables[33].rows),'table34 rows',len(d.tables[34].rows))
print('LITERAL ASTERISKS',sum(x.text.count('*') for x in d.paragraphs))
for n in [32,33,34]:
 print('TABLE',n)
 for r in d.tables[n].rows: print(' | '.join(c.text for c in r.cells))
