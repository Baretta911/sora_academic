from docx import Document
p=r'D:\Baretta\Skripsi\program\sora_academic\thesis_dump\Skripsi_123220204_BarettaP_format_template.docx'
d=Document(p)
for i,x in enumerate(d.paragraphs[860:]):
 if x.text.strip(): print(i+860,x.style.name,repr(x.text[:400]))
print('tables total',len(d.tables))
for i,t in enumerate(d.tables[-12:]): print('TABLE',len(d.tables)-12+i,'rows',len(t.rows),'cols',len(t.columns),'first',[[c.text[:80] for c in row.cells] for row in t.rows[:2]])
