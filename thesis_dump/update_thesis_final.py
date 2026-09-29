from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "Skripsi_123220204_BarettaP_revisi_model_Bab_II_III.docx"
OUT = ROOT / "Skripsi_123220204_BarettaP_revisi_final_data.docx"
d = Document(SRC)


def replace_paragraph(p, text):
    pPr = p._p.pPr
    for child in list(p._p):
        if child is not pPr:
            p._p.remove(child)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    p._p.append(r)


def clear_auto_number(p):
    pPr = p._p.pPr
    if pPr is not None:
        numPr = pPr.find(qn("w:numPr"))
        if numPr is not None:
            pPr.remove(numPr)


def set_heading(i, text):
    p = d.paragraphs[i]
    clear_auto_number(p)
    replace_paragraph(p, text)


def set_body(i, text):
    replace_paragraph(d.paragraphs[i], text)


def set_table(index, rows):
    t = d.tables[index]
    while len(t.rows) < len(rows):
        t.add_row()
    while len(t.rows) > len(rows):
        t._tbl.remove(t.rows[-1]._tr)
    assert len(rows[0]) == len(t.columns)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            p = cell.paragraphs[0]
            replace_paragraph(p, val)
            for extra in cell.paragraphs[1:]:
                replace_paragraph(extra, "")


# The chapter headings had inherited a single continuing automatic list from Chapter III.
set_heading(867, "BAB IV")
set_heading(868, "HASIL DAN PEMBAHASAN")
set_heading(870, "4.1 Hasil Kalibrasi Parameter")
set_heading(875, "4.2 Hasil Pengujian Optimasi Mingguan")
set_heading(880, "4.3 Hasil Pengujian Baseline Comparison")
set_heading(886, "4.4 Pembahasan")
set_heading(890, "BAB V")
set_heading(891, "PENUTUP")
set_heading(892, "5.1 Kesimpulan")
set_heading(898, "5.2 Saran")

set_body(869, "Bab ini menyajikan hasil eksperimen SORA pada dataset kuesioner final, yang memuat 63 baris dan 63 responden unik. Pengujian utama menggunakan seed 42, ukuran populasi 50, laju mutasi 0,10, dan laju crossover 0,80. Eksperimen utama dijalankan dengan opsi skip-calibration; karena itu parameter tersebut adalah konfigurasi default/fallback, bukan hasil kalibrasi yang diterapkan ulang pada pengujian final.")
set_body(871, "Kalibrasi grid search dilakukan terpisah pada lima subjek dengan 18 kombinasi ukuran populasi, laju mutasi, dan laju crossover. Nilai fitness rata-rata tertinggi adalah 2,917 pada populasi 100, mutasi 0,10, dan crossover 0,80. Hasil kalibrasi ini dilaporkan sebagai eksperimen terpisah. Metadata eksekusi utama mencatat populasi 50, mutasi 0,10, crossover 0,80, seed 42, dan parameter_source default_fallback.")
set_body(874, "Kombinasi dengan skor kalibrasi tertinggi tidak boleh disamakan dengan konfigurasi eksperimen utama. Pengujian final memakai konfigurasi fallback yang tercatat pada METADATA.json, sehingga hasil Bab IV di bawah menggambarkan keluaran konfigurasi tersebut.")
set_body(876, "Pengujian mingguan dijalankan pada seluruh 63 responden selama tujuh hari. Dengan demikian, terdapat 441 evaluasi harian dan seluruhnya berstatus optimized. Dataset tidur aktual digunakan sebagai baseline observasi, sedangkan sistem menghasilkan kandidat jadwal yang mempertahankan aktivitas wajib. Ringkasan metrik agregat disajikan pada Tabel 4.2.")
set_body(879, "Rata-rata fitness hasil optimasi adalah −2,2101 dan median fitness 6,9072. Rata-rata durasi tidur hasil optimasi 6,5045 jam per hari, sedangkan baseline observasi 6,5000 jam. Pada 399 dari 441 hari (90,48%), fitness hasil optimasi lebih tinggi daripada fitness baseline. Rata-rata fitness delta sebesar 156,7496; hasil negatif pada sebagian kasus menunjukkan bahwa kepadatan kendala dapat menghasilkan penalti besar. Angka rata-rata fitness yang negatif dipengaruhi sejumlah pencilan penalti, sehingga median juga dilaporkan.")
set_body(881, "Perbandingan baseline dilakukan pada 63 baris data hari Rabu dengan masukan yang sama untuk SORA, greedy, dan fixed schedule. Rata-rata fitness, fatigue akhir, serta durasi tidur disajikan pada Tabel 4.3.")
set_body(884, "Rata-rata fitness SORA adalah 6,2767, dibandingkan −52,8108 pada greedy dan −5630,9624 pada fixed. SORA menghasilkan fitness lebih tinggi daripada greedy pada 59 dari 63 baris dan daripada fixed pada 59 dari 63 baris; SORA menjadi metode dengan fitness tertinggi terhadap kedua baseline sekaligus pada 58 dari 63 baris. Rata-rata fatigue SORA 0,7354 lebih rendah daripada greedy 1,1828 dan fixed 0,8908. Durasi tidur rata-rata masing-masing adalah 6,5159 jam (SORA), 6,9048 jam (greedy), dan 7,4286 jam (fixed). Durasi yang lebih panjang tidak otomatis menghasilkan fitness lebih tinggi karena skor juga dipengaruhi pemulihan dan penalti kendala/regularitas.")
set_body(885, "Hasil tersebut mendukung keunggulan SORA pada skor fitness di sebagian besar data yang diuji, tetapi tidak berarti SORA selalu menghasilkan durasi tidur terpanjang. Perbandingan tetap dibatasi pada dataset, konfigurasi, seed, dan tiga metode yang dilaporkan.")
set_body(887, "Pada eksperimen mingguan, semua 63 responden memperoleh jadwal tujuh hari dan seluruh 441 evaluasi harian tercatat berhasil. Sebanyak 399 hari memiliki fitness lebih tinggi daripada baseline observasi. Rata-rata fitness −2,2101 dan median 6,9072 berbeda cukup jauh karena penalti yang sangat besar pada sejumlah jadwal; karena itu median dan rentang hasil perlu dipertimbangkan bersama rata-rata.")
set_body(888, "Pada perbandingan hari Rabu, SORA memiliki rata-rata fitness lebih tinggi daripada greedy dan fixed, serta rata-rata fatigue lebih rendah daripada kedua baseline. Durasi tidur rata-rata SORA sedikit lebih rendah daripada kedua pembanding, yang menunjukkan bahwa fungsi tujuan menilai beberapa komponen sekaligus dan tidak sekadar memaksimalkan jumlah jam tidur.")
set_body(889, "Hasil ini merupakan evaluasi komputasional atas fitness dan metrik turunannya. Penalti regularitas yang digunakan sistem bukan perhitungan Sleep Regularity Index (SRI) klinis penuh, dan keluaran SORA bukan diagnosis atau bukti dampak klinis. Generalisasi juga terbatas pada 63 responden dalam dataset ini serta satu konfigurasi seed dan parameter eksperimen utama.")

set_body(893, "Berdasarkan implementasi dan eksperimen pada data final, kesimpulan penelitian adalah sebagai berikut.")
set_body(894, "SORA berhasil mengoptimasi jadwal tidur tujuh hari untuk 63 responden unik. Seluruh 441 evaluasi harian berstatus optimized dengan seed 42 serta konfigurasi utama population size 50, mutation rate 0,10, dan crossover rate 0,80.")
set_body(895, "Pada optimasi mingguan, rata-rata fitness hasil optimasi sebesar −2,2101, median 6,9072, dan rata-rata durasi tidur 6,5045 jam per hari. Fitness hasil optimasi lebih tinggi daripada baseline observasi pada 399 dari 441 hari (90,48%).")
set_body(896, "Pada perbandingan 63 data hari Rabu, rata-rata fitness SORA sebesar 6,2767, lebih tinggi daripada greedy (−52,8108) dan fixed (−5630,9624). SORA mengungguli masing-masing baseline pada 59 dari 63 baris dan memiliki skor tertinggi dibanding keduanya sekaligus pada 58 baris.")
set_body(897, "Grid search terpisah pada lima subjek menghasilkan kombinasi terbaik populasi 100, mutasi 0,10, dan crossover 0,80 dengan fitness rata-rata 2,917. Konfigurasi ini berbeda dari konfigurasi utama yang benar-benar tercatat dijalankan pada data final, yaitu populasi 50, mutasi 0,10, dan crossover 0,80.")
set_body(899, "Berdasarkan hasil dan keterbatasan tersebut, pengembangan berikutnya dapat diarahkan pada hal-hal berikut.")
set_body(900, "Perbaiki mekanisme repair dan penalti adaptif untuk menangani subjek dengan kendala jadwal padat serta pencilan fitness yang negatif.")
set_body(901, "Ulangi eksperimen dengan beberapa random seed dan laporkan variasi hasil agar kestabilan metode dapat dinilai.")
set_body(902, "Bandingkan SORA dengan metode optimasi tambahan menggunakan data masukan dan metrik evaluasi yang sama.")
set_body(903, "Validasi sistem pada data pengguna nyata dan, jika tersedia, data wearable untuk menguji kesesuaian jadwal serta metrik secara longitudinal.")

# Replace stale Bab IV tables with values from the final pipeline artifacts and thesis notes.
set_table(33, [
    ["Metrik", "Nilai"],
    ["Baris data / responden unik", "63 / 63"],
    ["Evaluasi harian", "441 dari 441 berstatus optimized"],
    ["Fitness hasil optimasi (rata-rata / median)", "−2,2101 / 6,9072"],
    ["Fitness hasil optimasi (minimum / maksimum)", "−1346,9761 / 7,7997"],
    ["Fitness baseline observasi (rata-rata)", "−158,9597"],
    ["Fitness delta (rata-rata)", "156,7496"],
    ["Hari fitness optimasi lebih tinggi dari baseline", "399 dari 441 (90,48%)"],
    ["Durasi tidur optimasi / baseline (jam per hari)", "6,5045 / 6,5000"],
    ["Perubahan slot tidur (rata-rata)", "3,1474 slot (sekitar 1,57 jam)"],
    ["Fatigue akhir (rata-rata)", "0,6957"],
])
set_table(34, [
    ["Metrik", "SORA", "Greedy", "Fixed"],
    ["Rata-rata fitness", "6,2767", "−52,8108", "−5630,9624"],
    ["Rata-rata fatigue akhir", "0,7354", "1,1828", "0,8908"],
    ["Rata-rata durasi tidur (jam)", "6,5159", "6,9048", "7,4286"],
    ["Fitness SORA lebih tinggi", "—", "59 dari 63", "59 dari 63"],
    ["Fitness tertinggi simultan", "58 dari 63 data", "—", "—"],
])

# Correct the duplicate and shifted Chapter III equation labels introduced during insertion.
for table_index, label in [(13, "(3.4)"), (14, "(3.5)"), (22, "(3.6)"), (24, "(3.7)"), (29, "(3.9)")]:
    table = d.tables[table_index]
    for p in table.cell(0, 0).paragraphs:
        txt = p.text
        if "(3." in txt:
            replace_paragraph(p, txt.rsplit("(3.", 1)[0] + label)

# Remove unresolved cross-reference error strings wherever they appear; refresh fields on open.
for p in d.paragraphs:
    if "Error! Bookmark not defined." in p.text:
        replace_paragraph(p, p.text.replace("Error! Bookmark not defined.", ""))
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if "Error! Bookmark not defined." in p.text:
                    replace_paragraph(p, p.text.replace("Error! Bookmark not defined.", ""))
settings = d.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

d.save(OUT)
print(OUT)
