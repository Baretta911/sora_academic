import copy
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

DOCX = Path(r"D:\Baretta\Skripsi\penulisan\Skripsi_123220204_BarettaP.docx")
d = Document(str(DOCX))
paras = d.paragraphs


def insert_before_sectpr(el):
    body = d.element.body
    sect = body.find(qn("w:sectPr"))
    if sect is not None:
        body.insert(list(body).index(sect), el)
    else:
        body.append(el)


def set_text(p, text):
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    r = p.makeelement(qn("w:r"), {})
    t = p.makeelement(qn("w:t"), {})
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    p.append(r)


def cleanup_old_chapters():
    body = d.element.body
    anchor = None
    for p in d.paragraphs:
        if "Ringkasan Metodologi" in p.text:
            anchor = p._p
            break
    if anchor is None:
        return
    items = list(body)
    start = items.index(anchor) + 1
    removed = 0
    while start < len(body) and body[start].tag != qn("w:sectPr"):
        body.remove(body[start])
        removed += 1
    print("cleaned elements after anchor:", removed)


def clone_p(idx, text):
    el = copy.deepcopy(paras[idx]._p)
    set_text(el, text)
    insert_before_sectpr(el)
    return el


def add_table(rows, cols):
    t = d.add_table(rows=rows, cols=cols)
    t.style = d.tables[33].style
    insert_before_sectpr(t._tbl)
    return t


def fill(cell, text):
    set_text(cell.paragraphs[0]._p, text)


def find_idx(pred, start=0):
    for i in range(start, len(paras)):
        if pred(paras[i]):
            return i
    raise RuntimeError("template paragraph not found")


H1 = find_idx(lambda p: p.text.strip() == "BAB III")
H1T = find_idx(lambda p: p.text.strip() == "METODOLOGI PENELITIAN", H1 + 1)
H2 = find_idx(lambda p: p.text.strip() == "Jenis dan Rancangan Penelitian")
CAP = find_idx(lambda p: p.style.name == "Caption")
SMBR = find_idx(lambda p: p.text.strip().startswith("Sumber: Diolah"))
LST = find_idx(lambda p: p.style.name == "List Paragraph")
BODY = find_idx(lambda p: p.style.name == "Normal" and len(p.text.strip()) > 50, CAP + 1)
CLAUSE = find_idx(lambda p: "Dataset dalam penelitian ini dipisahkan" in p.text)

print("templates: bab=%(bab)s title=%(t)s h2=%(h2)s cap=%(cap)s sumber=%(smbr)s list=%(lst)s body=%(body)s clause=%(cl)s" % {
    "bab": H1, "t": H1T, "h2": H2, "cap": CAP, "smbr": SMBR, "lst": LST,
    "body": BODY, "cl": CLAUSE})


def fix_dataset_clause():
    p = paras[CLAUSE]
    txt = p.text.replace("terdiri atas 40 subjek", "terdiri atas 5 subjek")
    txt = txt.replace("terdiri atas 20 subjek", "terdiri atas 59 subjek")
    assert "40" not in txt and "20 subjek" not in txt, txt
    set_text(p._p, txt)


def kalibrasi_rows():
    rows = [
        ["50", "0,01", "0,70", "2,843"],
        ["50", "0,01", "0,80", "2,854"],
        ["50", "0,01", "0,90", "−263,881"],
        ["50", "0,05", "0,70", "2,904"],
        ["50", "0,05", "0,80", "2,904"],
        ["50", "0,05", "0,90", "−263,826"],
        ["50", "0,10", "0,70", "2,906"],
        ["50", "0,10", "0,80", "−263,824"],
        ["50", "0,10", "0,90", "2,915"],
        ["100", "0,01", "0,70", "−263,868"],
        ["100", "0,01", "0,80", "−263,851"],
        ["100", "0,01", "0,90", "2,884"],
        ["100", "0,05", "0,70", "2,905"],
        ["100", "0,05", "0,80", "−263,825"],
        ["100", "0,05", "0,90", "2,902"],
        ["100", "0,10", "0,70", "2,915"],
        ["100", "0,10", "0,80", "2,917"],
        ["100", "0,10", "0,90", "2,905"],
    ]
    t = add_table(19, 4)
    hdr = ["Ukuran Populasi", "Laju Mutasi", "Laju Crossover", "Rata-rata Fitness"]
    for j, h in enumerate(hdr):
        fill(t.cell(0, j), h)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            fill(t.cell(i, j), v)


def weekly_rows():
    return [
        ["Rata-rata fitness", "−0,755"],
        ["Median fitness", "6,819"],
        ["Fitness maksimum", "7,182"],
        ["Fitness minimum", "−186,438"],
        ["Subjek dengan fitness positif", "53 dari 59"],
        ["Rata-rata durasi tidur (jam)", "6,52"],
    ]


def baseline_rows():
    return [
        ["Rata-rata fitness", "6,369", "−33,173", "−5.942,307"],
        ["Rata-rata kelelahan akhir hari", "0,729", "1,183", "0,876"],
        ["Rata-rata durasi tidur (jam)", "6,51", "6,91", "7,57"],
    ]


def make_metric_table(caption, header, rows):
    clone_p(CAP, caption)
    t = add_table(len(rows) + 1, len(header))
    for j, h in enumerate(header):
        fill(t.cell(0, j), h)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            fill(t.cell(i, j), v)
    clone_p(SMBR, "Sumber: Diolah penulis")


def bab4():
    clone_p(H1, "BAB IV")
    clone_p(H1T, "HASIL DAN PEMBAHASAN")
    clone_p(BODY, "Bab ini menyajikan hasil pengujian sistem Sleep Optimization and Rest Arrangement (SORA), yaitu sistem optimizer yang menerapkan "
                  "Genetic Algorithm dan Hill Climbing untuk penyusunan jadwal tidur optimal, beserta pembahasannya. "
                  "Seluruh eksperimen dijalankan pada dataset kalibrasi yang terdiri atas lima subjek dan dataset "
                  "pengujian yang terdiri atas 59 subjek dengan menggunakan konfigurasi parameter terbaik hasil "
                  "kalibrasi serta random seed sebesar 42 agar hasil yang diperoleh dapat direproduksi. Pembahasan "
                  "difokuskan pada tiga skenario pengujian, yaitu kalibrasi parameter, pengujian optimasi mingguan "
                  "(weekly optimization test), dan pengujian perbandingan terhadap metode baseline (baseline comparison).")

    clone_p(H2, "Hasil Kalibrasi Parameter")
    clone_p(BODY, "Kalibrasi parameter dilakukan menggunakan metode grid search pada dataset kalibrasi yang terdiri "
                  "atas lima subjek. Grid search mengevaluasi 18 kombinasi parameter, yaitu kombinasi ukuran populasi "
                  "50 dan 100, laju mutasi 0,01; 0,05; dan 0,10, serta laju crossover 0,70; 0,80; dan 0,90. Setiap "
                  "kombinasi dijalankan selama 150 generasi Genetic Algorithm dilanjutkan dengan Hill Climbing pada "
                  "tiga hari representatif, yaitu Senin, Rabu, dan Sabtu. Hasil kalibrasi disajikan pada Tabel 4.1.")
    clone_p(CAP, "Tabel 4.1 Hasil Kalibrasi Parameter Menggunakan Grid Search")
    kalibrasi_rows()
    clone_p(SMBR, "Sumber: Diolah penulis")
    clone_p(BODY, "Berdasarkan Tabel 4.1 dapat diamati bahwa kombinasi parameter yang menghasilkan rata-rata "
                  "fitness tertinggi adalah kombinasi ukuran populasi 100, laju mutasi 0,10, dan laju crossover 0,80 "
                  "dengan nilai rata-rata fitness sebesar 2,917. Beberapa kombinasi lainnya menghasilkan nilai "
                  "rata-rata fitness yang mendekati nilai tersebut, antara lain kombinasi ukuran populasi 50, laju "
                  "mutasi 0,10, dan laju crossover 0,90 sebesar 2,915 serta kombinasi ukuran populasi 100, laju "
                  "mutasi 0,10, dan laju crossover 0,70 sebesar 2,915. Pada sebagian kombinasi, sistem menghasilkan "
                  "solusi dengan nilai fitness sangat rendah mendekati minus 263,8. Nilai tersebut diperoleh ketika "
                  "algoritma menghasilkan solusi yang melanggar kendala penjadwalan secara besar sehingga penalti "
                  "mendominasi nilai fitness. Hal ini menunjukkan bahwa hasil optimasi cukup peka terhadap kombinasi "
                  "parameter sehingga pemilihan parameter perlu dilakukan secara eksperimental. Berdasarkan hasil "
                  "tersebut, kombinasi ukuran populasi 100, laju mutasi 0,10, dan laju crossover 0,80 digunakan "
                  "sebagai konfigurasi default sistem dan digunakan pada seluruh skenario pengujian berikutnya.")

    clone_p(H2, "Hasil Pengujian Optimasi Mingguan")
    clone_p(BODY, "Pengujian optimasi mingguan dilakukan pada seluruh subjek dataset pengujian yang berjumlah 59 "
                  "subjek. Untuk setiap subjek, sistem menyusun jadwal tidur selama tujuh hari secara berurutan "
                  "dengan mempertimbangkan komponen kelelahan yang diwariskan antarhari (carry-over) sehingga solusi "
                  "yang dihasilkan memperhatikan keseimbangan antara aktivitas, tidur, dan istirahat pada level "
                  "mingguan. Statistik hasil pengujian optimasi mingguan disajikan pada Tabel 4.2.")
    make_metric_table("Tabel 4.2 Statistik Hasil Pengujian Optimasi Mingguan pada Seluruh Subjek Pengujian",
                      ["Metrik", "Nilai"], weekly_rows())
    clone_p(BODY, "Berdasarkan Tabel 4.2 dapat diketahui bahwa seluruh 59 subjek berhasil dioptimalkan dengan "
                  "status optimized pada ketujuh hari yang diuji. Median rata-rata fitness antar subjek sebesar "
                  "6,819 dengan nilai maksimum 7,182. Rata-rata durasi tidur yang dihasilkan sistem adalah 6,52 jam "
                  "per hari, mendekati target dasar tujuh jam dengan tetap memperhitungkan ketersediaan slot bebas "
                  "pada masing-masing hari. Lima subjek dengan rata-rata fitness tertinggi berturut-turut adalah "
                  "Fikri Fitriani (7,182), Dimas Baskoro (7,181), Dian Firmansyah (7,176), Eko Setiawan (7,167), dan "
                  "Yudha Susanto (7,162). Sebanyak 53 dari 59 subjek menghasilkan rata-rata fitness positif, "
                  "sedangkan enam subjek menghasilkan rata-rata fitness negatif yang akan dibahas lebih lanjut pada "
                  "subbab Pembahasan.")

    clone_p(H2, "Hasil Pengujian Baseline Comparison")
    clone_p(BODY, "Pengujian baseline comparison dilakukan pada hari Rabu sebagai hari representatif dengan "
                  "membandingkan sistem SORA terhadap dua metode pembanding, yaitu greedy schedule yang mengisi slot "
                  "bebas terpanjang dengan tidur dan fixed schedule yang menetapkan waktu tidur tetap pukul 22.00. "
                  "Ketiga metode dievaluasi menggunakan fungsi fitness yang sama. Hasil agregat dari pengujian "
                  "tersebut pada seluruh subjek pengujian disajikan pada Tabel 4.3.")
    make_metric_table("Tabel 4.3 Hasil Aggregat Pengujian Baseline Comparison",
                      ["Metrik", "SORA", "Greedy", "Fixed"], baseline_rows())
    clone_p(BODY, "Pada Tabel 4.3 terlihat bahwa sistem SORA menghasilkan rata-rata fitness sebesar 6,369. Nilai "
                  "tersebut jauh lebih tinggi dibandingkan metode greedy yang menghasilkan rata-rata fitness sebesar "
                  "−33,173 dan metode fixed yang menghasilkan rata-rata fitness sebesar −5.942,307. Nilai rata-rata "
                  "yang sangat rendah pada metode pembanding disebabkan oleh sejumlah subjek yang menghasilkan solusi "
                  "tidak layak, yaitu satu subjek pada metode greedy dan 17 subjek pada metode fixed. Hal tersebut "
                  "mengindikasikan bahwa penyusunan jadwal tidur dengan aturan sederhana tidak mampu memenuhi "
                  "kendala penjadwalan pada sebagian besar subjek.")
    clone_p(BODY, "Apabila hanya subjek yang menghasilkan solusi layak pada metode greedy yang diperhitungkan, "
                  "sistem SORA menghasilkan fitness yang lebih tinggi daripada metode greedy pada 54 dari 58 subjek. "
                  "Peningkatan fitness rata-rata sebesar 118,9% dengan peningkatan pada median sebesar 145,3%. Dari "
                  "sisi kelelahan akhir hari, sistem SORA menghasilkan rata-rata kelelahan sebesar 0,729, lebih "
                  "rendah dibandingkan metode greedy sebesar 1,183 sehingga terjadi penurunan kelelahan akhir hari "
                  "sebesar 38,4%. Hasil tersebut menunjukkan bahwa jadwal tidur yang disusun sistem SORA tidak hanya "
                  "unggul dalam nilai fitness, tetapi juga menurunkan kelelahan akhir hari secara lebih efektif "
                  "daripada skema penjadwalan sederhana.")

    clone_p(H2, "Pembahasan")
    clone_p(BODY, "Secara keseluruhan, hasil pengujian menunjukkan bahwa pendekatan Genetic Algorithm yang "
                  "dilanjutkan dengan Hill Climbing mampu menghasilkan jadwal tidur optimal pada seluruh subjek yang "
                  "diuji. Pada skenario optimasi mingguan, mayoritas subjek memperoleh rata-rata fitness positif "
                  "dengan median 6,819 dan durasi tidur rata-rata 6,52 jam per hari. Hasil ini sejalan dengan tujuan "
                  "sistem untuk menyusun jadwal yang menyeimbangkan tuntutan aktivitas, kebutuhan tidur, dan "
                  "kelelahan akhir hari tanpa harus mencapai durasi tidur maksimum secara kaku.")
    clone_p(BODY, "Hasil perbandingan terhadap metode baseline memperkuat keunggulan sistem. Jadwal tidur tetap "
                  "pukul 22.00 terbukti tidak layak pada 17 subjek karena benturan dengan aktivitas wajib yang "
                  "bersinggungan dengan waktu tidur tetap tersebut, sedangkan skema greedy menghasilkan jadwal yang"
                  "kurang memperhatikan hubungan antarhari dan kelelahan. Sistem SORA unggul pada 54 dari 58 subjek "
                  "dengan peningkatan fitness median sebesar 145,3% dan berhasil menurunkan kelelahan akhir hari "
                  "sebesar 38,4% dibandingkan metode greedy.")
    clone_p(BODY, "Meskipun demikian, terdapat enam subjek yang menghasilkan rata-rata fitness negatif pada "
                  "pengujian mingguan. Hal ini terutama berkaitan dengan kendala penjadwalan yang sangat ketat pada "
                  "subjek tersebut, misalnya banyaknya aktivitas wajib yang tetap dan sedikitnya slot bebas, sehingga "
                  "algoritma menghasilkan solusi dengan penalti yang besar. Kepekaan hasil terhadap kombinasi "
                  "parameter pada proses kalibrasi menunjukkan bahwa mekanisme perbaikan kelayakan dan penalti masih "
                  "dapat disempurnakan untuk menangani subjek dengan karakteristik kendala ekstrem.")

cleanup_old_chapters()

bab4()

clone_p(H1, "BAB V")
clone_p(H1T, "PENUTUP")
clone_p(H2, "Kesimpulan")
clone_p(BODY, "Berdasarkan hasil penelitian yang telah dilakukan, dapat ditarik kesimpulan sebagai berikut.")
for s in [
    "Sistem SORA berhasil diimplementasikan dengan menggunakan Genetic Algorithm dan Hill Climbing untuk menyusun "
    "jadwal tidur optimal. Kalibrasi parameter menggunakan grid search menetapkan kombinasi parameter terbaik pada "
    "ukuran populasi 100, laju mutasi 0,10, dan laju crossover 0,80 dengan rata-rata fitness sebesar 2,917 pada "
    "dataset kalibrasi.",
    "Pengujian optimasi mingguan pada 59 subjek menunjukkan bahwa sistem berhasil menyusun jadwal tidur selama "
    "tujuh hari dengan median rata-rata fitness 6,819 dan rata-rata durasi tidur 6,52 jam per hari. Sebanyak 53 "
    "dari 59 subjek menghasilkan rata-rata fitness positif.",
    "Hasil perbandingan terhadap metode baseline menunjukkan bahwa sistem SORA menghasilkan fitness yang lebih "
    "tinggi daripada metode greedy pada 54 dari 58 subjek dengan peningkatan rata-rata 118,9% dan peningkatan pada "
    "median sebesar 145,3%. Sistem juga menurunkan kelelahan akhir hari sebesar 38,4% dibandingkan metode greedy.",
    "Jadwal tidur tetap pukul 22.00 terbukti menghasilkan solusi tidak layak pada 17 subjek, sedangkan skema greedy "
    "tidak layak pada satu subjek, sehingga penyusunan jadwal dengan pendekatan optimasi terbukti lebih unggul "
    "daripada skema penjadwalan sederhana.",
]:
    clone_p(LST, s)
clone_p(H2, "Saran")
clone_p(BODY, "Berdasarkan hasil dan keterbatasan penelitian, terdapat beberapa saran untuk penelitian "
              "selanjutnya, antara lain sebagai berikut.")
for s in [
    "Perlu penyempurnaan mekanisme perbaikan kelayakan dan penalti adaptif agar dapat menangani subjek dengan "
    "kendala penjadwalan yang sangat ketat, mengingat masih terdapat enam subjek yang menghasilkan rata-rata "
    "fitness negatif pada pengujian mingguan.",
    "Perlu evaluasi hasil dalam beberapa kali pengulangan dengan random seed berbeda untuk mengurangi kepekaan hasil "
    "terhadap nilai seed tertentu.",
    "Perlu perbandingan kinerja algoritma dengan metode metaheuristik lain, misalnya Simulated Annealing atau "
    "algoritma genetik multiobjektif (NSGA-II), untuk memperkaya perbandingan pada penelitian lanjutan.",
    "Perlu pengujian sistem pada pengguna nyata serta integrasi dengan data wearable untuk memvalidasi pengaruh "
    "jadwal tidur yang dihasilkan terhadap kualitas tidur dan kelelahan secara longitudinal.",
]:
    clone_p(LST, s)

fix_dataset_clause()

before = len(paras)
d.save(str(DOCX))
d2 = Document(str(DOCX))
print("paragraphs before:", before, "after:", len(d2.paragraphs), "tables:", len(d2.tables))
print("clause475:", d2.paragraphs[475].text[:200])